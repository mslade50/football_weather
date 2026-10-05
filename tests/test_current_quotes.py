"""Elapsed build/upload time must not turn expired quotes into current prices."""
import csv
import json
from copy import deepcopy
from datetime import datetime, timedelta, timezone
from types import SimpleNamespace

import pytest

from pipeline import build
from pipeline.contracts import Game, GameLine
from pipeline.current_quotes import expire_card, expire_meta, expired, guard_legacy, publication_guard
from pipeline.outputs import json_out
from pipeline.outputs.legacy import LegacyRecord, cfb_game_label, nfl_game_label
from pipeline.run_context import RunContext
from utils.timeutil import date_label

NOW = datetime(2026, 10, 5, 2, 39, tzinfo=timezone.utc)
GID = "nfl:2026:4:a@h"


def card():
    return {"game_id": GID, "sport": "nfl", "away": {"name": "Atlanta Falcons"}, "home": {"name": "New Orleans Saints"},
            "kickoff_local": NOW.isoformat(), "odds": {"betcris": {"total": {"line": 48, "under": -110,
            "open_line": 47, "source_updated_at": "2026-10-05T02:23:02Z", "expires_at": "2026-10-05T02:28:02Z"}},
            "pinnacle": {"total": {"line": 49, "under": -111}}},
            "consensus": {"total_now": 48, "total_open": 47}, "fair": {"fair_total": 45, "edges": [{"book": "betcris"}]},
            "weather": {"rain_fg": None}, "total_prices": {"best": "betcris"}, "signal": {"label": "No Impact"}}


@pytest.mark.parametrize("value,expected", [(None, False), ("invalid", True), (NOW, False),
    (NOW - timedelta(microseconds=1), True), (NOW + timedelta(seconds=1), False), (NOW.replace(tzinfo=None), True)])
def test_expiry_boundary(value, expected):
    assert expired(value, NOW) is expected


def test_read_expiry_preserves_openers_and_source_clocks():
    old = card()
    new = expire_card(old, NOW)
    assert old == card()
    assert new["odds"]["betcris"]["total"]["open_line"] == 47
    assert new["odds"]["betcris"]["total"]["source_updated_at"] == "2026-10-05T02:23:02Z"
    assert "line" not in new["odds"]["betcris"]["total"]
    assert new["odds"]["pinnacle"] == old["odds"]["pinnacle"]
    assert new["consensus"]["total_now"] is None and new["consensus"]["total_open"] == 47
    assert new["fair"]["edges"] == [] and new["weather"]["rain_fg"] is None


def test_calculation_recomputes_consensus_fair_and_counts(monkeypatch):
    monkeypatch.setattr(RunContext, "now_utc", property(lambda self: NOW))
    game = Game(GID, "nfl", 2026, 4, NOW + timedelta(days=1), NOW + timedelta(days=1), "UTC", "h", "a", None)
    lines = [GameLine("nfl", GID, book, "total", side, -110, line=price,
                     scraped_at=NOW - timedelta(minutes=10), expires_at=expiry)
             for book, price, expiry in (("betcris", 48, NOW - timedelta(minutes=1)), ("pinnacle", 49, None))
             for side in ("over", "under")]
    openers = {"openers": {f"{GID}|total|under|betcris": {"line": 47}}}
    before = deepcopy(openers)
    odds = build.OddsResult(lines, {}, openers, {GID: lines}, [], {"betcris": 2, "pinnacle": 2}, scraped=list(lines))
    impact = SimpleNamespace(gs_fg_pct=-5, away_fg_pct=0, rain_c=0)
    record = LegacyRecord("nfl", "a", "h", game.kickoff_local, game_id=GID)
    res = build.SportResult("nfl", [record], [], [card()], [game], {}, {}, {}, {GID: impact}, odds)
    ctx = RunContext("nfl", git_sha="test")
    ctx.counts["betcris"] = {"nfl": 2, "nfl.total": 2}
    build.refresh_expired_quotes(ctx, res)
    assert res.cards[0]["consensus"]["total_now"] == 49
    assert "betcris" not in res.cards[0]["odds"]
    assert all(e["book"] != "betcris" for e in res.cards[0]["fair"]["edges"])
    assert record.odds["total_proj"] == 49 and ctx.counts["betcris"]["nfl"] == 0
    assert openers == before and len(res.odds.scraped) == 4
    build.refresh_expired_quotes(ctx, res)
    assert len(ctx.degradations) == 1


def test_calculation_with_no_live_books_keeps_displayed_historical_opener(monkeypatch):
    monkeypatch.setattr(RunContext, "now_utc", property(lambda self: NOW))
    game = Game(GID, "nfl", 2026, 4, NOW + timedelta(days=1), NOW + timedelta(days=1), "UTC", "h", "a", None)
    line = GameLine("nfl", GID, "betcris", "total", "under", -110, line=48,
                    scraped_at=NOW - timedelta(minutes=10), expires_at=NOW - timedelta(minutes=1))
    odds = build.OddsResult([line], {}, {"openers": {}}, {GID: [line]}, [], {"betcris": 1}, scraped=[line])
    record = LegacyRecord("nfl", "a", "h", game.kickoff_local, game_id=GID)
    record.odds = {"total_now": 48, "total_open": 47, "under_open": -110}
    res = build.SportResult("nfl", [record], [], [card()], [game], {}, {}, {}, {}, odds)
    build.refresh_expired_quotes(RunContext("nfl", git_sha="test"), res)
    assert record.odds["total_now"] is None
    assert record.odds["total_open"] == 47 and record.odds["under_open"] == -110
    assert res.cards[0]["consensus"]["total_now"] is None and not res.cards[0]["fair"]["edges"]
    assert odds.scraped == [line]


def test_publication_after_slow_upload_removes_current_prices_and_keeps_history(tmp_path):
    directory = tmp_path / "board"
    directory.mkdir()
    original = card()
    meta = {"run_id": "r", "books": {"betcris": {"count": 2, "status": "green"}}, "counts": {"betcris": {"nfl": 2}},
            "quote_expiries": [{"book": "betcris", "sport": "nfl", "market": "total", "count": 2,
                                "expires_at": "2026-10-05T02:28:02Z"}]}
    json_out.dump_json(directory / "meta.json", meta)
    json_out.dump_json(directory / "status.json", {"meta": {}, "counts": meta["counts"], "books": meta["books"]})
    (tmp_path / "state").mkdir()
    json_out.dump_json(tmp_path / "state/status.json", {})
    json_out.dump_json(directory / "games_nfl.json", {"meta": {}, "games": [original]})
    history = {"series": {"opener": [47]}}
    json_out.dump_json(directory / "history.json", history)
    path = tmp_path / "nfl_weather.csv"
    with path.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, ["Game", "Date", "Total_now", "Total_open"])
        w.writeheader()
        w.writerow({"Game": nfl_game_label(original["away"]["name"], original["home"]["name"]), "Date": date_label(NOW),
                    "Total_now": 48, "Total_open": 47})
    publication_guard(directory, NOW)
    assert json.loads((directory / "meta.json").read_text())["books"]["betcris"]["status"] == "red"
    assert json.loads((tmp_path / "state/status.json").read_text())["books"]["betcris"]["status"] == "red"
    assert json.loads((directory / "history.json").read_text()) == history
    with path.open(encoding="utf-8") as f:
        row = next(csv.DictReader(f))
    assert row["Total_now"] == "" and row["Total_open"] == "47"
    assert expire_meta(expire_meta(meta, NOW), NOW) == expire_meta(meta, NOW)


def test_excel_expiry_preserves_opener_weather_and_other_games(tmp_path):
    from openpyxl import Workbook, load_workbook
    original = card()
    original["sport"] = "cfb"
    label = cfb_game_label(original["away"]["name"], original["home"]["name"])
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Game", "Date", "Total_now", "Total_open", "Temp_fg"])
    sheet.append([label, date_label(NOW), 48, 47, 75])
    sheet.append(["another game", date_label(NOW), 51, 50, 76])
    path = tmp_path / "cfb_weather.xlsx"
    workbook.save(path)
    workbook.close()
    guard_legacy(tmp_path, [expire_card(original, NOW)])
    workbook = load_workbook(path)
    assert list(workbook.active.values)[1:] == [(label, date_label(NOW), None, 47, 75),
                                              ("another game", date_label(NOW), 51, 50, 76)]
    workbook.close()


@pytest.mark.parametrize("remaining,withheld", [(30, True), (-1, True), (120, False), (121, False)])
def test_publication_reserves_remaining_upload_time_without_shifting_clocks(tmp_path, remaining, withheld):
    directory = tmp_path / "board"
    directory.mkdir()
    original = card()
    expiry = NOW + timedelta(seconds=remaining)
    quote = original["odds"]["betcris"]["total"]
    quote["expires_at"] = expiry.isoformat()
    quote["source_updated_at"] = (expiry - timedelta(seconds=300)).isoformat()
    meta = {"books": {"betcris": {"count": 4, "status": "green"}},
            "counts": {"betcris": {"nfl": 2, "cfb": 2}}, "quote_expiries": [
                {"book": "betcris", "sport": "nfl", "market": "total", "count": 2, "expires_at": expiry.isoformat()},
                {"book": "betcris", "sport": "cfb", "market": "total", "count": 2,
                 "expires_at": (NOW + timedelta(minutes=30)).isoformat()}]}
    json_out.dump_json(directory / "meta.json", meta)
    json_out.dump_json(directory / "games_nfl.json", {"meta": {}, "games": [original]})
    publication_guard(directory, NOW, window_seconds=120)
    checked = json.loads((directory / "games_nfl.json").read_text())["games"][0]
    after = json.loads((directory / "meta.json").read_text())
    assert checked["odds"]["betcris"]["total"].get("expired", False) is withheld
    assert checked["odds"]["betcris"]["total"]["expires_at"] == quote["expires_at"]
    assert checked["odds"]["betcris"]["total"]["source_updated_at"] == quote["source_updated_at"]
    assert checked["odds"]["betcris"]["total"]["open_line"] == quote["open_line"]
    assert after["publication_guard_at"] == NOW.isoformat() and after["publication_window_seconds"] == 120
    assert after["counts"]["betcris"]["cfb"] == 2
    if withheld:
        assert after["counts"]["betcris"]["nfl"] == 0 and after["books"]["betcris"]["status"] == "amber"
        assert after["degradations"][-1]["ts"] == NOW.isoformat()
        assert "publication window" in after["degradations"][-1]["reason"]
        assert checked["fair"]["fair_total"] is None
    else:
        assert after["books"]["betcris"]["count"] == 4 and not after.get("degradations")


def test_compact_board_carries_earliest_current_quote_expiry():
    original = card()
    original.update(season=2026, week=4, kickoff_utc=NOW.isoformat(), date_label="Mon", time_label="2:39", neutral=False)
    original["home"]["short"], original["away"]["short"] = "NO", "ATL"
    assert json_out.table_row(original)["quote_expires_at"] == "2026-10-05T02:28:02Z"
    assert json_out.table_row(expire_card(original, NOW))["quote_expires_at"] is None


@pytest.mark.parametrize("stamp,expected", [("2026-10-05T09:17:00+00:00", "2026-10-05T14:17:00Z"),
    ("2026-10-05T20:18:00+00:00", "2026-10-05T22:17:00Z"),
    ("2026-10-04T16:18:00+00:00", "2026-10-04T16:47:00Z"),
    ("2026-10-03T11:18:00+00:00", "2026-10-03T12:17:00Z"),
    ("2026-11-01T16:18:00+00:00", "2026-11-01T16:47:00Z")])
def test_cadence_matches_existing_utc_triggers_across_dst(stamp, expected):
    assert json_out.next_run_eta(datetime.fromisoformat(stamp), {}) == expected


def test_next_backstop_is_pinned_to_workflow_schedule():
    import re
    from pathlib import Path
    crons = re.findall(r"cron: '([^']+)'", (Path(__file__).parents[1] / ".github/workflows/pipeline.yml").read_text())
    def matches(field, value):
        for part in field.split(","):
            base, _, step = part.partition("/")
            low, _, high = base.partition("-")
            if base == "*":
                return True
            end = int(high or low)
            if int(low) <= value <= end and (value - int(low)) % int(step or 1) == 0:
                return True
        return False
    for day in range(7):
        now = NOW.replace(hour=8, minute=0) + timedelta(days=day)
        for _ in range(18):
            expected = now.replace(second=0, microsecond=0) + timedelta(minutes=1)
            while not any(matches(c.split()[0], expected.minute) and matches(c.split()[1], expected.hour)
                          and matches(c.split()[4], (expected.weekday() + 1) % 7) for c in crons):
                expected += timedelta(minutes=1)
            assert json_out.next_backstop(now) == expected
            now = expected
