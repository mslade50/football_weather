"""Expire current quotes without changing their clocks or historical openers."""
from __future__ import annotations

import argparse
import json
from collections import Counter
from copy import deepcopy
from datetime import datetime, timedelta, timezone
from pathlib import Path


def expired(value, now: datetime) -> bool:
    if value is None:
        return False
    try:
        expiry = value if isinstance(value, datetime) else datetime.fromisoformat(value.replace("Z", "+00:00"))
        return expiry.tzinfo is None or now > expiry
    except (ValueError, TypeError, AttributeError):
        return True


def _datetime(value) -> datetime | None:
    if isinstance(value, datetime):
        return value
    if not isinstance(value, str) or not value:
        return None
    try:
        parsed = datetime.fromisoformat(value.replace("Z", "+00:00"))
    except ValueError:
        return None
    return parsed if parsed.tzinfo is not None else parsed.replace(tzinfo=timezone.utc)


def _card_lines(card: dict) -> list | None:
    """Rebuild the serialized main lines needed by the existing fair model.

    Exchange quotes omit their raw probability from the display block; the fair
    edge block retains that exact, already-devigged probability. If it is absent,
    return None rather than replacing a precise exchange quote with rounded
    American odds.
    """
    sport, game_id = card.get("sport"), card.get("game_id")
    if sport not in ("nfl", "cfb") or not game_id:
        return None
    from pipeline.contracts import GameLine
    from pipeline.model.fair import EXCHANGE_BOOKS

    edge_probs = {}
    for edge in (card.get("fair") or {}).get("edges") or []:
        if not isinstance(edge, dict):
            continue
        key = (edge.get("book"), edge.get("market"), edge.get("side"))
        edge_probs[key] = edge.get("vigfree_prob")

    lines = []
    for book, markets in (card.get("odds") or {}).items():
        if not isinstance(markets, dict):
            continue
        for market, quote in markets.items():
            if market not in ("spread", "total", "ml") or not isinstance(quote, dict) or quote.get("expired"):
                continue
            sides = []
            if market == "spread":
                home_line = quote.get("home_line")
                if home_line is None:
                    continue
                if quote.get("home_odds") is not None:
                    sides.append(("home", home_line, quote["home_odds"]))
                if quote.get("away_odds") is not None:
                    sides.append(("away", -home_line, quote["away_odds"]))
            elif market == "total":
                line = quote.get("line")
                if line is None:
                    continue
                for side in ("over", "under"):
                    if quote.get(side) is not None:
                        sides.append((side, line, quote[side]))
            else:
                for side in ("home", "away"):
                    if quote.get(side) is not None:
                        sides.append((side, None, quote[side]))

            for side, line, odds in sides:
                if not isinstance(odds, int) or isinstance(odds, bool):
                    return None
                prob_raw = None
                if book in EXCHANGE_BOOKS and market in ("spread", "total"):
                    prob_raw = edge_probs.get((book, market, side))
                    if not isinstance(prob_raw, (int, float)) or not 0 < prob_raw < 1:
                        return None
                lines.append(GameLine(
                    sport=sport, game_id=game_id, book=book, market=market, side=side,
                    odds=odds, line=line, prob_raw=prob_raw, is_main=True,
                    scraped_at=_datetime(quote.get("updated_at")),
                    source_updated_at=_datetime(quote.get("source_updated_at")),
                    expires_at=_datetime(quote.get("expires_at")),
                ))
    return lines


def _recompute_card(card: dict) -> bool:
    """Refresh comparisons from surviving serialized quotes; preserve all openers."""
    lines = _card_lines(card)
    impact = card.get("impact") or {}
    v1 = impact.get("v1")
    if lines is None or not isinstance(v1, dict):
        return False

    from pipeline.model import fair as fair_mod
    from pipeline.outputs import json_out

    sport, game_id = card["sport"], card["game_id"]
    wx = card.get("weather") or {}
    stadium = card.get("stadium") or {}
    v1_components = v1.get("components") or {}
    args = {
        "wind_vol_fc": wx.get("wind_vol_fc"),
        "wind_vol_static": stadium.get("wind_vol_static"),
        "model_disagreement": wx.get("model_disagreement"),
        "lead_hours": wx.get("lead_hours"),
    }
    fair_v1 = fair_mod.evaluate_game(
        sport, game_id, lines, v1.get("gs_fg_pct"), v1.get("away_fg_pct"),
        rain_c=v1_components.get("rain"), **args,
    )
    fair_v2 = None
    v2 = impact.get("v2")
    if isinstance(v2, dict) and any(v2.get(key) is not None for key in ("gs_fg_pct", "away_fg_pct")):
        v2_components = v2.get("components") or {}
        fair_v2 = fair_mod.evaluate_game(
            sport, game_id, lines, v2.get("gs_fg_pct"), v2.get("away_fg_pct"),
            rain_c=v2_components.get("rain"), model_version="v2", **args,
        )
    selected = fair_v2 if impact.get("model_version") == "v2" and fair_v2 is not None else fair_v1

    consensus = card.get("consensus") or {}
    sp, total = selected.spread, selected.total
    sp_now, total_now = sp.line, total.line
    consensus.update({
        "spread_now": sp_now,
        "total_now": total_now,
        "spread_src": sp.src or ("fallback" if sp_now is not None else None),
        "move_s": sp_now - consensus["spread_open"] if sp_now is not None and consensus.get("spread_open") is not None else None,
        "move_t": total_now - consensus["total_open"] if total_now is not None and consensus.get("total_open") is not None else None,
        "ref_book": total.ref_book or sp.ref_book,
        "n_books": max(sp.n_books, total.n_books),
        "thin": max(sp.n_books, total.n_books) < 2,
    })
    card["consensus"] = consensus

    legacy = None
    legacy_fn = getattr(fair_mod, "legacy_derived", None)
    if callable(legacy_fn):
        ref = "fanduel" if sport == "cfb" else "betonline"
        ref_total = (card.get("odds") or {}).get(ref, {}).get("total") or {}
        current_total = ref_total.get("line")
        if current_total is None:
            current_total = total_now
        legacy = legacy_fn(total_now, sp_now, current_total, v1.get("gs_fg_pct"), v1.get("away_fg_pct"))
    card["fair"] = json_out.fair_block(selected, legacy, fair_v2)
    card["total_prices"] = json_out.compare_totals(sport, lines, selected)

    # Execution references must not outlive the only current total quote for a book.
    if "execution_markets" in card:
        card["execution_markets"] = [
            ref for ref in card["execution_markets"]
            if (card.get("odds") or {}).get(ref.get("book"), {}).get("total", {}).get("line") is not None
            and not (card.get("odds") or {}).get(ref.get("book"), {}).get("total", {}).get("expired")
        ]
    return True


def expire_card(card: dict, now: datetime) -> dict:
    card = deepcopy(card)
    removed = []
    for book, markets in card.get("odds", {}).items():
        for market, quote in markets.items():
            if not expired(quote.get("expires_at"), now):
                continue
            removed.append(f"{book}/{market}")
            markets[market] = {k: v for k, v in quote.items()
                               if k.startswith("open") or k in ("expires_at", "updated_at", "source_updated_at")}
            markets[market]["expired"] = True
    if removed:
        card["expired_markets"] = removed
        # Fresh book lines remain useful. Rebuild consensus/fair/edges from the
        # surviving main quotes using the same model functions as the full build.
        # Older/incomplete cards fail closed, while preserving historical openers.
        try:
            recomputed = _recompute_card(card)
        except Exception:  # noqa: BLE001 - expire safely if a serialized card cannot be rehydrated
            # Incomplete or malformed cards must never retain comparisons that
            # might still include the expired quote.
            recomputed = False
        if not recomputed:
            card["consensus"] = {**card.get("consensus", {}), **dict.fromkeys(
                ("spread_now", "total_now", "move_s", "move_t", "spread_src", "ref_book")), "n_books": 0,
                "thin": True}
            card["fair"] = {k: [] if k == "edges" else None for k in card.get("fair", {})}
            card["total_prices"] = None
    return card


def expire_meta(meta: dict, now: datetime) -> dict:
    meta = deepcopy(meta)
    if "quote_expiries" not in meta:
        return meta
    live = []
    removed = Counter()
    for group in meta.get("quote_expiries", []):
        if expired(group["expires_at"], now):
            removed[(group["book"], group["sport"], group["market"])] += group["count"]
        else:
            live.append(group)
    meta["quote_expiries"] = live
    for (book, sport, market), count in removed.items():
        counts = meta.get("counts", {}).get(book, {})
        for key in (sport, f"{sport}.{market}"):
            if key in counts:
                counts[key] = max(0, counts[key] - count)
        status = meta.get("books", {}).get(book)
        if status is not None:
            status["count"] = max(0, (status.get("count") or 0) - count)
            status["status"] = "amber" if status["count"] else "red"
            status["reason"] = "Expired quotes excluded from current prices"
    if removed:
        meta.setdefault("degradations", []).append({"component": "odds.expiry", "severity": "warn",
            "reason": f"{sum(removed.values())} expired quotes excluded; source timestamps unchanged",
            "run_id": meta.get("run_id"), "ts": now.isoformat()})
    return meta


def publication_guard(directory: Path, now: datetime, *, window_seconds: int = 0) -> None:
    """Recheck after raw uploads, immediately before uploading current board JSON."""
    cutoff = now + timedelta(seconds=window_seconds)
    original = json.loads((directory / "meta.json").read_text(encoding="utf-8"))
    meta = expire_meta(original, cutoff)
    if window_seconds:
        meta["publication_guard_at"] = now.isoformat()
        meta["publication_window_seconds"] = window_seconds
        for degradation in meta.get("degradations", [])[len(original.get("degradations", [])):]:
            degradation["ts"] = now.isoformat()
            degradation["reason"] = degradation["reason"].replace("expired quotes excluded", "quotes excluded for expiry within the publication window")
        for book, status in meta.get("books", {}).items():
            if status != original.get("books", {}).get(book):
                status["reason"] = "Quotes expiring within the publication window excluded from current prices"
    cards = []
    for sport in ("nfl", "cfb"):
        path = directory / f"games_{sport}.json"
        if not path.exists():
            continue
        payload = json.loads(path.read_text(encoding="utf-8"))
        payload["games"] = [expire_card(c, cutoff) for c in payload["games"]]
        payload["meta"] = {**payload.get("meta", {}), "degradations": meta.get("degradations", [])}
        cards.extend(payload["games"])
        path.write_text(json.dumps(payload, allow_nan=False), encoding="utf-8")
    from pipeline.outputs.json_out import table_row
    guard_legacy(directory.parent, cards)
    board = directory / "board.json"
    if board.exists():
        payload = json.loads(board.read_text(encoding="utf-8"))
        payload["rows"] = [table_row(c) for c in cards]
        payload["meta"] = {**payload.get("meta", {}), "degradations": meta.get("degradations", [])}
        board.write_text(json.dumps(payload, allow_nan=False), encoding="utf-8")
    status = directory / "status.json"
    if status.exists():
        payload = json.loads(status.read_text(encoding="utf-8"))
        for key in ("degradations", "books", "counts"):
            payload[key] = meta.get(key, {})
        payload["meta"] = {**payload.get("meta", {}), "degradations": meta.get("degradations", [])}
        payload["quote_expiries"] = meta.get("quote_expiries", [])
        status.write_text(json.dumps(payload, allow_nan=False), encoding="utf-8")
        state_status = directory.parent / "state/status.json"
        if state_status.exists():
            state_status.write_text(json.dumps(payload, allow_nan=False), encoding="utf-8")
    (directory / "meta.json").write_text(json.dumps(meta, allow_nan=False), encoding="utf-8")


def guard_legacy(directory: Path, cards: list[dict]) -> None:
    """Clear current comparisons in legacy snapshots; opening/history fields stay."""
    import csv

    from pipeline.outputs.legacy import cfb_game_label, nfl_game_label
    from utils.timeutil import date_label
    labels = {(c["sport"], (nfl_game_label if c["sport"] == "nfl" else cfb_game_label)(c["away"]["name"], c["home"]["name"]),
               date_label(datetime.fromisoformat(c["kickoff_local"]))) for c in cards if c.get("expired_markets")}
    fields = {"Spread_now", "Odds_now", "Total_now", "Under_now", "FD_now", "Odds_n", "Current", "Spread",
              "Total_proj", "Move_t", "Move_s", "My_total", "Edge", "My_spread", "Edge_s"}
    path = directory / "nfl_weather.csv"
    if path.exists() and any(s == "nfl" for s, _, _ in labels):
        with path.open(encoding="utf-8", newline="") as source:
            reader = csv.DictReader(source)
            headers, rows = reader.fieldnames, list(reader)
        for row in rows:
            if ("nfl", row.get("Game"), row.get("Date")) in labels:
                for key in fields & row.keys():
                    row[key] = ""
        with path.open("w", encoding="utf-8", newline="") as output:
            writer = csv.DictWriter(output, headers)
            writer.writeheader()
            writer.writerows(rows)
    path = directory / "cfb_weather.xlsx"
    if path.exists() and any(s == "cfb" for s, _, _ in labels):
        from openpyxl import load_workbook
        workbook = load_workbook(path)
        for sheet in workbook:
            headers = {cell.value: cell.column for cell in sheet[1]}
            if "Game" not in headers or "Date" not in headers:
                continue
            for row in sheet.iter_rows(min_row=2):
                if ("cfb", row[headers["Game"] - 1].value, row[headers["Date"] - 1].value) in labels:
                    for key in fields & headers.keys():
                        row[headers[key] - 1].value = None
        workbook.save(path)
        workbook.close()


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--board-dir", type=Path, required=True)
    parser.add_argument("--publication-window-seconds", type=int, default=0,
                        help="Withhold quotes expiring during the remaining board/state uploads")
    args = parser.parse_args()
    if not 0 <= args.publication_window_seconds <= 600:
        parser.error("publication window must be between 0 and 600 seconds")
    publication_guard(args.board_dir, datetime.now(timezone.utc), window_seconds=args.publication_window_seconds)
